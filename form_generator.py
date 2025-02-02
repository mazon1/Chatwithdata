import streamlit as st
import pandas as pd
import os
import pdfplumber
import pytesseract
import sqlite3
import fitz  # PyMuPDF for alternative PDF text extraction
from pdf2image import convert_from_path
from PIL import Image
import docx
from fpdf import FPDF
import google.generativeai as genai
from sentence_transformers import SentenceTransformer, util

# Configure Google Gemini API
GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY", st.secrets.get("GOOGLE_API_KEY"))
genai.configure(api_key=GOOGLE_API_KEY)

# Load Embedding Model for RAG
embedding_model = SentenceTransformer('all-MiniLM-L6-v2')

# Initialize SQLite Database
conn = sqlite3.connect("uploaded_data.db")
cursor = conn.cursor()

cursor.execute("""
CREATE TABLE IF NOT EXISTS documents (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    file_name TEXT,
    extracted_text TEXT,
    embedding BLOB
)
""")
conn.commit()

# Function to Extract Text from PDFs (Using pdfplumber & PyMuPDF as fallback)
def extract_text_from_pdf(pdf_file):
    text = ""
    try:
        with pdfplumber.open(pdf_file) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
    except Exception:
        doc = fitz.open(stream=pdf_file.read(), filetype="pdf")
        text = "\n".join([page.get_text("text") for page in doc])
    return text if text else "No text found."

# Function to Extract Text from Images (OCR)
def extract_text_from_image(image_file):
    image = Image.open(image_file)
    return pytesseract.image_to_string(image)

# Function to Extract Text from Word Documents
def extract_text_from_docx(doc_file):
    doc = docx.Document(doc_file)
    return "\n".join([para.text for para in doc.paragraphs])

# Function to Extract Data from Excel
def extract_data_from_excel(excel_file):
    df = pd.read_excel(excel_file)
    return df.to_csv(index=False)  

# Store extracted content in SQLite
def save_to_db(file_name, extracted_text):
    embedding = embedding_model.encode(extracted_text, convert_to_tensor=True).tolist()
    cursor.execute("INSERT INTO documents (file_name, extracted_text, embedding) VALUES (?, ?, ?)",
                   (file_name, extracted_text, str(embedding)))
    conn.commit()

# Function to Generate a Response with Gemini
def generate_response(prompt):
    try:
        model = genai.GenerativeModel('gemini-pro')
        response = model.generate_content(prompt)
        return response.text  
    except Exception as e:
        return f"Error generating response: {e}"

# Function to Retrieve Relevant Content for RAG
def retrieve_relevant_content(query):
    query_embedding = embedding_model.encode(query, convert_to_tensor=True)
    
    # Fetch embeddings from DB
    cursor.execute("SELECT file_name, extracted_text, embedding FROM documents")
    results = cursor.fetchall()
    
    best_match = None
    best_score = 0
    for file_name, text, embedding in results:
        stored_embedding = eval(embedding)  # Convert back to list
        score = util.pytorch_cos_sim(query_embedding, stored_embedding).item()
        if score > best_score:
            best_match = text
            best_score = score
    
    return best_match if best_match else "No relevant data found."

# Function to Populate a Template
def generate_report(template_text, extracted_data):
    report_text = template_text
    for key, value in extracted_data.items():
        report_text = report_text.replace(f"{{{{{key}}}}}", value)
    return report_text

# Function to Export as PDF
def export_to_pdf(report_text, filename="generated_report.pdf"):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    for line in report_text.split("\n"):
        try:
            pdf.cell(200, 10, txt=line.encode("latin-1", "ignore").decode("latin-1"), ln=True)
        except UnicodeEncodeError:
            pdf.cell(200, 10, txt="(Encoding Error: Some characters could not be displayed)", ln=True)

    pdf.output(filename)
    return filename

# Streamlit UI
st.title("📄 Multi-File Data Extraction, AI Insights & Report Generator")

uploaded_files = st.file_uploader("Upload multiple documents", type=["pdf", "docx", "jpg", "png", "xlsx"], accept_multiple_files=True)

if uploaded_files:
    extracted_data = {}
    
    for uploaded_file in uploaded_files:
        file_type = uploaded_file.type
        
        if "pdf" in file_type:
            extracted_text = extract_text_from_pdf(uploaded_file)
        elif "image" in file_type or "png" in file_type or "jpg" in file_type:
            extracted_text = extract_text_from_image(uploaded_file)
        elif "word" in file_type or "docx" in file_type:
            extracted_text = extract_text_from_docx(uploaded_file)
        elif "spreadsheet" in file_type or "xlsx" in file_type:
            extracted_text = extract_data_from_excel(uploaded_file)
        else:
            extracted_text = "Unsupported file format."

        save_to_db(uploaded_file.name, extracted_text)
        extracted_data[uploaded_file.name] = extracted_text

    st.success("✅ Documents uploaded and processed!")

# Template Upload & Report Generation
st.header("📑 Generate Report from Template")

template_file = st.file_uploader("Upload a report template (Word, Text, or Excel file)", type=["docx", "txt", "xlsx"])
if template_file:
    if "docx" in template_file.type:
        template_text = extract_text_from_docx(template_file)
    elif "xlsx" in template_file.type:
        template_text = extract_data_from_excel(template_file)
    else:
        template_text = template_file.read().decode()

    st.text_area("📜 Template Preview:", template_text, height=200)

    doc_options = st.selectbox("Select document data to use:", extracted_data.keys() if extracted_data else [])
    if doc_options:
        selected_text = extracted_data[doc_options]

        report_text = generate_report(template_text, {"EXTRACTED_DATA": selected_text})
        st.text_area("📄 Generated Report:", report_text, height=300)

        if st.button("Export as PDF"):
            pdf_filename = export_to_pdf(report_text)
            st.success("📂 Report generated successfully!")
            st.download_button("Download Report", open(pdf_filename, "rb"), file_name=pdf_filename, mime="application/pdf")

# AI Chatbot Section
st.header("💬 AI Chatbot for Insights")

user_input = st.text_input("Ask a question about the extracted data:")
if st.button("Send"):
    retrieved_text = retrieve_relevant_content(user_input)
    response = generate_response(f"Using the following data, answer concisely: {retrieved_text}")
    st.write(f"🤖 AI: {response}")
